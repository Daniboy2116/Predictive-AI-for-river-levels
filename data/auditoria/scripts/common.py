"""Carga y normalización de todas las fuentes de niveles del Orinoco."""
import os
import re
import datetime as dt
import calendar
import pandas as pd
import numpy as np
import openpyxl

AUDITORIA = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
# Los 11 archivos originales (extraídos de "data new.rar"); la carpeta no se versiona
DATA = os.path.join(os.path.dirname(AUDITORIA), "data new")
SALIDA = os.path.join(AUDITORIA, "salida")
os.makedirs(SALIDA, exist_ok=True)

STATIONS = ["Ayacucho", "Caicara", "Ciudad_Bolivar", "Palua"]
STATION_LABELS = {
    "ayacucho": "Ayacucho",
    "caicara": "Caicara",
    "ciudad_bolivar": "Ciudad_Bolivar",
    "ciudad bolivar": "Ciudad_Bolivar",
    "ciudad bolívar": "Ciudad_Bolivar",
    "palua": "Palua",
    "palúa": "Palua",
    "puerto ayacucho": "Ayacucho",
    "caicara el mini mini": "Caicara",
}

MONTHS = {m: i for i, m in enumerate(
    ["ENE", "FEB", "MAR", "ABR", "MAY", "JUN", "JUL", "AGO", "SEP", "OCT", "NOV", "DIC"], 1)}


def norm_station(name):
    key = str(name).strip().lower()
    return STATION_LABELS.get(key, name)


def _clean_numeric(v):
    if v is None:
        return np.nan
    if isinstance(v, (int, float, np.integer, np.floating)):
        return float(v) if not (isinstance(v, float) and np.isnan(v)) else np.nan
    s = str(v).strip().replace(",", ".")
    if s == "" or s.lower() in ("nan", "none", "null", "-"):
        return np.nan
    try:
        return float(s)
    except ValueError:
        return np.nan


def load_wide_csv(path, source_name):
    """CSV/XLSX en formato ancho: una columna por año, 366 filas = día del año.
    Convención observada: fila 1 = 1-ene; fila 60 = 29-feb (solo bisiestos, NaN si no);
    fila 366 = 31-dic en bisiestos y en no bisiestos es el 31-dic tras saltar el hueco."""
    raw = pd.read_csv(path, header=0)
    years = [int(c) for c in raw.columns]
    rows = []
    for yi, year in enumerate(years):
        leap = calendar.isleap(year)
        for idx in range(len(raw)):
            if (not leap) and idx == 59:
                continue
            if leap:
                date = dt.date(year, 1, 1) + dt.timedelta(days=idx)
            else:
                doy = idx + 1 if idx < 59 else idx
                date = dt.date(year, 1, 1) + dt.timedelta(days=doy - 1)
            val = _clean_numeric(raw.iloc[idx, yi])
            if pd.isna(val):
                continue
            rows.append((date, source_name, val))
    return pd.DataFrame(rows, columns=["Fecha", "Fuente", "Valor"])


def load_wide_xlsx_plain(path, source_name):
    """XLSX en formato ancho sin encabezado de años (ayacucho_raw.xlsx)."""
    raw = pd.read_excel(path, header=None)
    years = [int(c) for c in raw.iloc[0].tolist()]
    rows = []
    for yi, year in enumerate(years):
        leap = calendar.isleap(year)
        for idx in range(1, len(raw)):
            if (not leap) and idx - 1 == 59:
                continue
            if leap:
                date = dt.date(year, 1, 1) + dt.timedelta(days=idx - 1)
            else:
                doy = idx if idx <= 59 else idx - 1
                date = dt.date(year, 1, 1) + dt.timedelta(days=doy - 1)
            val = _clean_numeric(raw.iloc[idx, yi])
            if pd.isna(val):
                continue
            rows.append((date, source_name, val))
    return pd.DataFrame(rows, columns=["Fecha", "Fuente", "Valor"])


def load_datos_sheet(ws, source_name):
    """Hoja del libro DATOS RIO ORINOCO: DIAS, mes, día, años..."""
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    years = [int(y) for y in hdr[3:] if y is not None]
    # Algunas hojas (Ayacucho, Palúa) dejan vacía la columna DIAS en la fila
    # del 29-feb; no debe descartarse esa fila si tiene día y valor.
    data = [r for r in rows[1:] if r[0] is not None or r[2] is not None]
    month = None
    out = []
    for r in data:
        if r[1] is not None:
            m = str(r[1]).strip().upper()
            if m in MONTHS:
                month = MONTHS[m]
        day = r[2]
        if month is None or day is None:
            continue
        for j, year in enumerate(years):
            try:
                date = dt.date(year, month, int(day))
            except ValueError:
                continue
            val = _clean_numeric(r[3 + j])
            if pd.isna(val):
                continue
            out.append((date, source_name, val))
    return pd.DataFrame(out, columns=["Fecha", "Fuente", "Valor"])


def load_entrenar(ws):
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h).strip() if h is not None else "" for h in rows[0]]
    idx = {}
    for i, h in enumerate(hdr[:5]):
        if i == 0:
            continue
        idx[norm_station(h)] = i
    out = []
    for r in rows[1:]:
        date = r[0]
        if date is None:
            continue
        if isinstance(date, dt.datetime):
            date = date.date()
        elif isinstance(date, str):
            try:
                date = pd.to_datetime(date).date()
            except Exception:
                continue
        else:
            continue
        for st, j in idx.items():
            val = _clean_numeric(r[j])
            if not pd.isna(val):
                out.append((date, "DATOS:Entrenar", st, val))
    return pd.DataFrame(out, columns=["Fecha", "Fuente", "Estacion", "Valor"])


def load_dataset_csv(path, source_name):
    df = pd.read_csv(path)
    df.columns = [c.strip().lower() for c in df.columns]
    out = []
    for _, r in df.iterrows():
        try:
            date = pd.to_datetime(r["fecha"], errors="coerce")
        except Exception:
            continue
        if pd.isna(date):
            continue
        date = date.date()
        for c in ["ayacucho", "caicara", "ciudad_bolivar", "palua"]:
            val = _clean_numeric(r[c])
            if not pd.isna(val):
                out.append((date, source_name, norm_station(c), val))
    return pd.DataFrame(out, columns=["Fecha", "Fuente", "Estacion", "Valor"])


def load_dataset_xlsx(path, source_name):
    df = pd.read_excel(path)
    df.columns = [str(c).strip().lower() for c in df.columns]
    out = []
    for _, r in df.iterrows():
        date = r["fecha"]
        if isinstance(date, str):
            date = pd.to_datetime(date, errors="coerce")
        if pd.isna(date):
            continue
        date = date.date() if hasattr(date, "date") else date
        for c in ["ayacucho", "caicara", "ciudad_bolivar", "palua"]:
            val = _clean_numeric(r[c])
            if not pd.isna(val):
                out.append((date, source_name, norm_station(c), val))
    return pd.DataFrame(out, columns=["Fecha", "Fuente", "Estacion", "Valor"])


def load_all():
    """Devuelve dict fuente -> DataFrame largo (Fecha, Fuente, Estacion, Valor)."""
    sources = {}

    # los CSV anchos no tienen columna de estación: loader propio
    def wide_csv(path, station, name):
        raw = pd.read_csv(path, header=0)
        years = [int(c) for c in raw.columns]
        rows = []
        for yi, year in enumerate(years):
            leap = calendar.isleap(year)
            for idx in range(len(raw)):
                if (not leap) and idx == 59:
                    continue
                if leap:
                    date = dt.date(year, 1, 1) + dt.timedelta(days=idx)
                else:
                    doy = idx + 1 if idx < 59 else idx
                    date = dt.date(year, 1, 1) + dt.timedelta(days=doy - 1)
                val = _clean_numeric(raw.iloc[idx, yi])
                if not pd.isna(val):
                    rows.append((date, name, station, val))
        return pd.DataFrame(rows, columns=["Fecha", "Fuente", "Estacion", "Valor"])

    sources["ayacucho_raw.csv"] = wide_csv(
        os.path.join(DATA, "ayacucho_raw.csv"), "Ayacucho", "ayacucho_raw.csv")
    sources["caicara_raw.csv"] = wide_csv(
        os.path.join(DATA, "caicara_raw.csv"), "Caicara", "caicara_raw.csv")
    sources["ciudad_bolivar_raw.csv"] = wide_csv(
        os.path.join(DATA, "ciudad_bolivar_raw.csv"), "Ciudad_Bolivar", "ciudad_bolivar_raw.csv")
    sources["palua_raw.csv"] = wide_csv(
        os.path.join(DATA, "palua_raw.csv"), "Palua", "palua_raw.csv")

    sources["ayacucho_raw.xlsx"] = load_wide_xlsx_plain(
        os.path.join(DATA, "ayacucho_raw.xlsx"), "ayacucho_raw.xlsx")
    sources["ayacucho_raw.xlsx"]["Estacion"] = "Ayacucho"
    sources["ayacucho_raw.xlsx"] = sources["ayacucho_raw.xlsx"][
        ["Fecha", "Fuente", "Estacion", "Valor"]]

    wb = openpyxl.load_workbook(os.path.join(DATA, "DATOS RIO ORINOCO.xlsx"),
                                read_only=True, data_only=True)
    sheet_station = {
        "Puerto Ayacucho": "Ayacucho",
        "Caicara": "Caicara",
        "Ciudad Bolivar": "Ciudad_Bolivar",
        "Palua": "Palua",
    }
    for sh, st in sheet_station.items():
        name = f"DATOS:{sh}"
        df = load_datos_sheet(wb[sh], name)
        df["Estacion"] = st
        sources[name] = df[["Fecha", "Fuente", "Estacion", "Valor"]]
    ent = load_entrenar(wb["Entrenar"])
    sources["DATOS:Entrenar"] = ent
    wb.close()

    for fname, name in [
        ("dataset_combinado.csv", "dataset_combinado.csv"),
        ("dataset_hasta_2006.csv", "dataset_hasta_2006.csv"),
        ("dataset_imputado_simpleml.csv", "dataset_imputado_simpleml.csv"),
    ]:
        sources[name] = load_dataset_csv(os.path.join(DATA, fname), name)

    for fname, name in [
        ("dataset_imputado_simpleml.xlsx", "dataset_imputado_simpleml.xlsx"),
        ("dataset-orinoco-imputado (1).xlsx", "dataset-orinoco-imputado (1).xlsx"),
    ]:
        sources[name] = load_dataset_xlsx(os.path.join(DATA, fname), name)

    for k, v in sources.items():
        v["Fecha"] = pd.to_datetime(v["Fecha"]).dt.date
        v["Valor"] = v["Valor"].astype(float)
        v.drop_duplicates(subset=["Fecha", "Estacion"], keep="first", inplace=True)
        v.sort_values(["Estacion", "Fecha"], inplace=True)
        v.reset_index(drop=True, inplace=True)
    return sources


def series(sources, station):
    """dict fuente -> Serie indexada por fecha para una estación."""
    out = {}
    for name, df in sources.items():
        d = df[df["Estacion"] == station]
        if len(d):
            out[name] = pd.Series(d["Valor"].values, index=pd.to_datetime(d["Fecha"]))
            out[name] = out[name][~out[name].index.duplicated(keep="first")].sort_index()
    return out
