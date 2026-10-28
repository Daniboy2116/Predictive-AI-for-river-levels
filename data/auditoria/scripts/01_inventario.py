"""Paso 1: inventario de archivos, hojas, columnas, rangos, NaN y estadísticas."""
import os
import sys
import json
import pandas as pd
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import common

src = common.load_all()

FILE_META = {
    "ayacucho_raw.csv": ("CSV ancho (años en columnas, 366 filas)", 366, 52),
    "caicara_raw.csv": ("CSV ancho (años en columnas, 366 filas)", 366, 52),
    "ciudad_bolivar_raw.csv": ("CSV ancho (años en columnas, 366 filas)", 366, 52),
    "palua_raw.csv": ("CSV ancho (años en columnas, 366 filas)", 366, 52),
    "ayacucho_raw.xlsx": ("XLSX ancho (años en columnas, 366 filas)", 366, 52),
    "DATOS:Puerto Ayacucho": ("Hoja XLSX ancha con mes/día y años", 366, 66),
    "DATOS:Caicara": ("Hoja XLSX ancha con mes/día y años", 366, 61),
    "DATOS:Ciudad Bolivar": ("Hoja XLSX ancha con mes/día y años", 366, 94),
    "DATOS:Palua": ("Hoja XLSX ancha con mes/día y años", 366, 86),
    "DATOS:Entrenar": ("Hoja XLSX larga (fecha + 4 estaciones + Pred)", 5934, 40),
    "dataset_combinado.csv": ("CSV largo (fecha + 4 estaciones)", 18683, 5),
    "dataset_hasta_2006.csv": ("CSV largo (fecha + 4 estaciones)", 12053, 5),
    "dataset_imputado_simpleml.csv": ("CSV largo (fecha + 4 estaciones)", 18683, 5),
    "dataset_imputado_simpleml.xlsx": ("XLSX largo (fecha + 4 estaciones)", 18683, 5),
    "dataset-orinoco-imputado (1).xlsx": ("XLSX largo (1 hoja) (fecha + 4 estaciones)", 18683, 5),
}

files_rows = []
station_rows = []
for name, df in src.items():
    desc, nrows, ncols = FILE_META[name]
    stations = sorted(df["Estacion"].unique())
    files_rows.append({
        "archivo": name,
        "formato": desc,
        "filas": nrows,
        "columnas": ncols,
        "estaciones": ", ".join(stations),
        "n_estaciones": len(stations),
        "fecha_min": df["Fecha"].min().isoformat(),
        "fecha_max": df["Fecha"].max().isoformat(),
        "n_valores": len(df),
    })
    for st in stations:
        d = df[df["Estacion"] == st]
        fmin, fmax = d["Fecha"].min(), d["Fecha"].max()
        expected = (fmax - fmin).days + 1
        station_rows.append({
            "archivo": name,
            "estacion": st,
            "fecha_min": fmin.isoformat(),
            "fecha_max": fmax.isoformat(),
            "n": len(d),
            "dias_esperados": expected,
            "pct_nan": round(100 * (expected - len(d)) / expected, 3),
            "min": round(d["Valor"].min(), 3),
            "max": round(d["Valor"].max(), 3),
            "mediana": round(d["Valor"].median(), 3),
            "unidad_probable": "m",
        })

fdf = pd.DataFrame(files_rows)
sdf = pd.DataFrame(station_rows)
fdf.to_csv(os.path.join(common.SALIDA, "inventario_archivos.csv"), index=False)
sdf.to_csv(os.path.join(common.SALIDA, "inventario_estaciones.csv"), index=False)

print(fdf.to_string(index=False))
print()
print(sdf.to_string(index=False))

# Chequeos adicionales: hojas y dimensiones reales del libro
wb = openpyxl.load_workbook(os.path.join(common.DATA, "DATOS RIO ORINOCO.xlsx"),
                            read_only=True, data_only=True)
sheet_dims = {ws.title: (ws.max_row, ws.max_column) for ws in wb.worksheets}
wb.close()
with open(os.path.join(common.SALIDA, "inventario_libro.json"), "w", encoding="utf-8") as f:
    json.dump(sheet_dims, f, ensure_ascii=False, indent=2)
print("\nDimensiones de hojas del libro:", sheet_dims)

# Comprobación de duplicados exactos entre datasets de 4 estaciones
comb = common.series(src, "Ayacucho")["dataset_combinado.csv"]
imp_csv = common.series(src, "Ayacucho")["dataset_imputado_simpleml.csv"]
imp_xlsx = common.series(src, "Ayacucho")["dataset_imputado_simpleml.xlsx"]
imp_1 = common.series(src, "Ayacucho")["dataset-orinoco-imputado (1).xlsx"]
print("\nIgualdad imputado csv vs xlsx:", imp_csv.equals(imp_xlsx))
print("Igualdad imputado csv vs (1).xlsx:", imp_csv.equals(imp_1))
